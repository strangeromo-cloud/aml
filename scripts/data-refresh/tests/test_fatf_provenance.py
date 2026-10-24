"""A retained snapshot must never turn a failed FATF attempt into fresh data."""
from __future__ import annotations

import copy
import importlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
builder = importlib.import_module("scripts.data-refresh.build_legal_workbook")
verifier = importlib.import_module("scripts.data-refresh.verify_fatf")

BLACK = ["Iran", "North Korea", "Myanmar"]
GREY = ["Angola", "Bolivia", "Bulgaria", "Cameroon", "Haiti", "Kenya", "Lebanon", "Monaco"]
ROWS = [{"status": status, "country": country, "publication_date": "19 June 2026"}
        for status, countries in [("Call for Action", BLACK), ("Increased Monitoring", GREY)]
        for country in countries]
BASELINE = {"listDate": "2026-06-19", "rows": ROWS}
ABSENT = object()


class FatfProvenanceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.snapshots = self.root / "_snapshots"
        self.snapshots.mkdir()
        self.manifest = self.root / "_manifest.json"
        for module in [builder, verifier]:
            for name, value in [("OUTPUT_DIR", self.root), ("SNAPSHOT_DIR", self.snapshots),
                                ("MANIFEST_FILE", self.manifest)]:
                mock = patch.object(module, name, value)
                mock.start()
                self.addCleanup(mock.stop)
        for name in ["CPI_OVERRIDE", "OFFSHORE_OVERRIDE"]:
            mock = patch.object(builder, name, self.root / name)
            mock.start()
            self.addCleanup(mock.stop)
        mock = patch.object(builder, "country_code", return_value="XX")
        mock.start()
        self.addCleanup(mock.stop)
        self.baseline = patch.object(verifier, "load_baseline", return_value=copy.deepcopy(BASELINE)).start()
        self.addCleanup(patch.stopall)

    def prepare(self, status="error", rows=ROWS):
        sources = {}
        if status is not ABSENT:
            sources["fatf-jurisdictions"] = {"status": status}
            if status == "error":
                sources["fatf-jurisdictions"]["error"] = "HTTP 403; Cloudflare verification required"
        self.manifest.write_text(json.dumps({"sources": sources}))
        path = self.snapshots / "fatf-jurisdictions.json"
        if rows is not ABSENT:
            path.write_text(json.dumps(rows))

    def build(self):
        out = self.root / "attachment.xlsx"
        report = builder.build(out)["fatf"]
        workbook = load_workbook(out)
        self.addCleanup(workbook.close)
        return report, workbook["FATF"]

    def assert_baseline_failure(self, report, sheet):
        self.assertEqual(report["source"], "baseline")
        self.assertFalse(report["ok"])
        self.assertTrue(report["seeded"])
        self.assertIsNone(report["agreesWithFetch"])
        self.assertIsNone(report["fetchDate"])
        self.assertIn("使用法务基线（非本次抓取）", sheet["A1"].value)
        self.assertNotIn("与本次官方抓取一致", sheet["A1"].value)
        self.assertIn("2026-06-19", sheet["A1"].value)
        self.assertEqual({row[1].value for row in sheet.iter_rows(min_row=3)}, set(BLACK + GREY))

    def test_failed_fetch_with_matching_old_snapshot_uses_explicit_baseline(self):
        self.prepare()
        report, sheet = self.build()
        self.assert_baseline_failure(report, sheet)
        self.assertTrue(sheet["A1"].value.startswith("⚠ 本次抓取失败，使用法务基线"))
        self.assertTrue(sheet["A1"].alignment.wrap_text)
        self.assertEqual(sheet.row_dimensions[1].height, 54)
        _, payload, error = verifier.source_fresh_fetch()
        self.assertIsNone(payload)
        self.assertIn("HTTP 403", error)

    def test_failed_fetch_ignores_newer_dated_retained_snapshot(self):
        rows = copy.deepcopy(ROWS)
        for row in rows:
            row["publication_date"] = "23 October 2026"
        rows[0]["country"] = "Stale cache sentinel"
        self.prepare(rows=rows)
        report, sheet = self.build()
        self.assert_baseline_failure(report, sheet)
        self.assertNotIn("Stale cache sentinel", str(list(sheet.values)))

    def test_missing_or_non_success_status_cannot_adopt_a_snapshot(self):
        for status in [ABSENT, None, "skipped", "pending"]:
            with self.subTest(status=status):
                self.prepare(status=status)
                report, sheet = self.build()
                self.assert_baseline_failure(report, sheet)
                self.assertIn("本次未确认抓取成功", sheet["A1"].value)
                self.assertIsNone(verifier.source_fresh_fetch()[1])

    def test_failed_fetch_without_baseline_does_not_publish_old_snapshot(self):
        self.prepare()
        self.baseline.return_value = None
        report, sheet = self.build()
        self.assertEqual(report["source"], "unavailable")
        self.assertFalse(report["ok"])
        self.assertFalse(report["seeded"])
        self.assertEqual(report["records"], 0)
        self.assertEqual(sheet["B3"].value, "抓取失败，无数据")

    def test_successful_matching_fetch_retains_agreement_behavior(self):
        self.prepare(status="success")
        report, sheet = self.build()
        self.assertTrue(report["ok"])
        self.assertTrue(report["agreesWithFetch"])
        self.assertIn("与本次官方抓取一致", sheet["A1"].value)
        self.assertNotIn("非本次抓取", sheet["A1"].value)
        _, payload, error = verifier.source_fresh_fetch()
        self.assertIsNone(error)
        self.assertEqual(payload["rows"], ROWS)

    def test_successful_newer_fetch_can_still_supersede_baseline(self):
        rows = copy.deepcopy(ROWS)
        for row in rows:
            row["publication_date"] = "23 October 2026"
        self.prepare(status="success", rows=rows)
        report, _ = self.build()
        self.assertTrue(report["ok"])
        self.assertFalse(report["seeded"])
        self.assertEqual(report["source"], "fetch")

    def test_success_without_snapshot_is_not_verifiable(self):
        self.prepare(status="success", rows=ABSENT)
        report, sheet = self.build()
        self.assert_baseline_failure(report, sheet)
        self.assertIn("未产出可用快照", report["reason"])
        self.assertIsNone(verifier.source_fresh_fetch()[1])

    def test_success_with_partial_snapshot_is_rejected(self):
        self.prepare(status="success", rows=[r for r in ROWS if r["status"] == "Increased Monitoring"])
        report, sheet = self.build()
        self.assert_baseline_failure(report, sheet)
        self.assertIsNone(verifier.source_fresh_fetch()[1])

    def test_success_with_corrupt_snapshot_is_not_fresh_data(self):
        self.prepare(status="success")
        (self.snapshots / "fatf-jurisdictions.json").write_text("not-json")
        report, sheet = self.build()
        self.assert_baseline_failure(report, sheet)
        self.assertIsNone(verifier.source_fresh_fetch()[1])

    def test_validator_exception_fails_closed(self):
        self.prepare(status="success")
        with patch.object(verifier, "validate_rows", side_effect=ValueError("invalid schema")):
            report, sheet = self.build()
            self.assert_baseline_failure(report, sheet)
            self.assertIsNone(verifier.source_fresh_fetch()[1])

    def test_missing_or_corrupt_manifest_never_confirms_fresh_data(self):
        self.prepare(status="success")
        for contents in [None, "not-json"]:
            with self.subTest(contents=contents):
                if contents is None:
                    self.manifest.unlink(missing_ok=True)
                else:
                    self.manifest.write_text(contents)
                report, sheet = self.build()
                self.assert_baseline_failure(report, sheet)
                self.assertIsNone(verifier.source_fresh_fetch()[1])

    def test_invalid_manifest_structure_never_confirms_fresh_data(self):
        self.prepare(status="success")
        for manifest in [[], {"sources": None}, {"sources": []},
                         {"sources": {"fatf-jurisdictions": ["success"]}}]:
            with self.subTest(manifest=manifest):
                self.manifest.write_text(json.dumps(manifest))
                report, sheet = self.build()
                self.assert_baseline_failure(report, sheet)
                self.assertIsNone(verifier.source_fresh_fetch()[1])


if __name__ == "__main__":
    unittest.main()
