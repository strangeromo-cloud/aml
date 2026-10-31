"""EU FSF XML with session cookies, bounded retry and full-list validation.

Webgate can set cookies on a redirect back to the same URL. Keep one in-memory
cookie jar for this fetch; urllib's normal redirect limits still apply. CSV has
different name/remark/country semantics and is not a transparent XML replacement.
"""
from __future__ import annotations

import http.client
import http.cookiejar
import logging
import ssl
import time
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
from typing import Any

from openpyxl import load_workbook

from ..common import DEFAULT_HEADERS, OUTPUT_DIR, Fetcher, FetchError, load_manifest

XML_URL = ("https://webgate.ec.europa.eu/fsd/fsf/public/files/"
           "xmlFullSanctionsList_1_1/content?token=dG9rZW4tMjAxNw")
MIN_ENTITIES = 3000
MIN_RETAINED_RATIO = 0.8
RETRY_DELAY_SECONDS = 120


def download_xml(url: str, opener) -> bytes:
    """Retry transient transport failures once after a two-minute cooldown.

    Reuse the opener so cookies received even by a failed redirect remain useful.
    Do not include URLs, query values or cookie values in diagnostics.
    """
    for attempt in range(2):
        retryable = True
        try:
            req = urllib.request.Request(url, headers={
                **DEFAULT_HEADERS, "Accept": "application/xml,text/xml;q=0.9"})
            with opener.open(req, timeout=120) as response:
                return response.read()
        except urllib.error.HTTPError as exc:
            reason = f"HTTP {exc.code}"
            retryable = 300 <= exc.code < 400 or exc.code == 429 or 500 <= exc.code < 600
            exc.close()
        except (urllib.error.URLError, TimeoutError, ConnectionError, ssl.SSLError,
                http.client.HTTPException) as exc:
            reason = type(exc).__name__
        if not retryable or attempt == 1:
            raise FetchError(f"EU XML download failed ({reason}; {attempt + 1} attempt(s))") from None
        logging.getLogger("data-refresh").warning(
            "EU XML download failed (%s); retrying once in %s seconds with the same cookie session",
            reason, RETRY_DELAY_SECONDS)
        time.sleep(RETRY_DELAY_SECONDS)
    raise AssertionError("unreachable")


def _parse_xml(raw: bytes) -> list[dict[str, Any]]:
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        raise FetchError(f"EU XML parse failed: {exc}") from None
    for el in root.iter():
        if isinstance(el.tag, str) and "}" in el.tag:
            el.tag = el.tag.split("}", 1)[1]
    if root.tag != "export":
        raise FetchError("EU XML has an unexpected root; expected the full FSF export")

    rows: list[dict[str, Any]] = []
    for entity in root.findall("./sanctionEntity"):
        sub = entity.find("subjectType")
        kind_code = sub.get("code") if sub is not None else ""
        kind_label = "Person" if (kind_code or "").upper().startswith("P") else "Entity"
        names = []
        primary_name = ""
        for alias in entity.findall(".//nameAlias"):
            whole_name = alias.attrib.get("wholeName") or ""
            if not whole_name.strip():
                parts = [alias.attrib.get(k, "").strip() for k in ("firstName", "middleName", "lastName")]
                whole_name = " ".join(p for p in parts if p)
            if whole_name:
                if not primary_name:
                    primary_name = whole_name
                else:
                    names.append(whole_name)
        countries = []
        for citizen in entity.findall(".//citizenship"):
            country = citizen.attrib.get("countryDescription") or citizen.attrib.get("country") or ""
            if country:
                countries.append(country)
        for address in entity.findall(".//address"):
            country = address.attrib.get("countryDescription") or ""
            if country:
                countries.append(country)
        reg = entity.find("regulation")
        regulation = (reg.attrib.get("publicationTitle") or reg.attrib.get("numberTitle") or "") if reg is not None else ""
        listed_on = next((r.attrib["entryIntoForceDate"] for r in entity.findall("regulation")
                          if r.attrib.get("entryIntoForceDate")), "")
        rows.append({
            "kind": kind_label,
            # Preserve the established column contract; this is the EU reference,
            # not the numeric sanctionEntity.logicalId used in the CSV subrecords.
            "logical_id": entity.attrib.get("euReferenceNumber", ""),
            "listed_on": listed_on, "name": primary_name,
            "aliases": " ; ".join(names), "countries": " ; ".join(sorted(set(countries))),
            "regulation": regulation, "remarks": (entity.findtext("remark") or "").strip(),
        })
    return rows


def _validate_rows(rows: list[dict[str, Any]], previous_count: int | None = None) -> None:
    if len(rows) < MIN_ENTITIES:
        raise FetchError(f"EU XML contains only {len(rows)} entities (minimum {MIN_ENTITIES}); keeping the stored list")
    seen = set()
    for row in rows:
        identity = str(row.get("logical_id") or "").strip()
        if not identity or identity in seen:
            raise FetchError("EU XML has missing or duplicate EU reference numbers; keeping the stored list")
        seen.add(identity)
        if not str(row.get("name") or "").strip():
            raise FetchError("EU XML has an entity without a name; keeping the stored list")
    if previous_count and len(rows) < previous_count * MIN_RETAINED_RATIO:
        raise FetchError(f"EU entity count fell from {previous_count} to {len(rows)} "
                         f"(more than {1 - MIN_RETAINED_RATIO:.0%}); manual review required, keeping the stored list")


class EuConsolidatedFetcher(Fetcher):
    id = "eu-consolidated"
    name = "EU Consolidated Financial Sanctions List"
    url = XML_URL
    out_filename = "eu-consolidated.xlsx"

    def columns(self):
        return [
            ("kind", "Kind", 12),
            ("logical_id", "Logical ID", 12),
            ("listed_on", "Listed On", 14),
            ("name", "Name", 50),
            ("aliases", "Aliases", 60),
            ("countries", "Countries", 30),
            ("regulation", "Regulation", 26),
            ("remarks", "Remarks", 60),
        ]

    def extra_about_rows(self):
        return [("Fetch strategy", "XML with an in-memory cookie session"),
                ("Entity validation", "Unique EU references, non-empty names, count floor and previous-list comparison")]

    def _previous_count(self) -> int | None:
        def count(value):
            return value if isinstance(value, int) and not isinstance(value, bool) and value > 0 else None
        previous = (load_manifest().get("sources") or {}).get(self.id) or {}
        if isinstance(previous, dict):
            known = count(previous.get("records")) or count(previous.get("previousRecords"))
            if known:
                return known
        # Error manifest entries may no longer carry a record count. The last
        # published workbook is retained across failed downloads and remains a
        # comparison baseline, including after several consecutive failures.
        path = OUTPUT_DIR / self.out_filename
        if path.exists():
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    known = count(wb["About"]["B5"].value)
                    if known and wb["Data"].max_row == known + 1:
                        return known
                finally:
                    wb.close()
            except Exception:
                logging.getLogger("data-refresh").warning("EU stored workbook count unavailable; using the absolute entity floor")
        return None

    def fetch(self) -> list[dict[str, Any]]:
        cookies = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cookies))
        previous = self._previous_count()
        rows = _parse_xml(download_xml(self.url, opener))
        _validate_rows(rows, previous)
        return rows
